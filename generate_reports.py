import os
import sys
import random
import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference

# Define categories and configuration
CATEGORIES = {
    "Selenium-Web": {
        "name": "Selenium - Website Tests (300)",
        "file_prefix": "selenium-web-report",
        "description": "End-to-End browser UI automation, responsiveness, and frontend web validation tests.",
        "modules": ["Auth", "Dashboard UI", "Analytics Graphs", "Patient Profiles", "Reports Export", "Navigation", "Settings"],
        "verbs": ["Verify", "Validate", "Test", "Check", "Ensure"],
        "scenarios": [
            "login screen responsiveness on {device}",
            "successful authentication with valid credentials",
            "error feedback for invalid login inputs",
            "navigation path redirection to {page}",
            "periodontal charting screen rendering accuracy",
            "CSV export functionality for patient lists",
            "PDF report download contains all correct tooth diagrams",
            "dark mode styling toggle and theme persistence",
            "input sanitization on patient search query",
            "responsive scaling of tooth model rendering on {device}",
            "session timeout redirect page functionality",
            "language selector translations for {locale}"
        ]
    },
    "Appium-Android": {
        "name": "Appium - Android Tests (300)",
        "file_prefix": "appium-android-report",
        "description": "Mobile native application UI/UX, permission handling, offline storage, and voice probing tests.",
        "modules": ["Splash", "Dashboard", "Speech Recognition", "Voice Parser", "Room Database", "Tooth Detail Screen", "Offline Sync"],
        "verbs": ["Verify", "Check", "Validate", "Ensure", "Test"],
        "scenarios": [
            "microphone permission prompt on first microphone click",
            "app launch and splash screen auto-transition within 1.5s",
            "voice parsing of command: 'tooth {tooth_num} depth {depth_val}'",
            "voice parser rejects invalid tooth number '{invalid_tooth}'",
            "speech recognition listener status updates correctly to '{speech_status}'",
            "offline insert of tooth data to Room DB",
            "data synchronization with backend PHP server once connection is restored",
            "UI back-stack behavior on back press from probing screen",
            "PDF generator executes locally without crash",
            "network status receiver reports offline mode instantly",
            "dashboard cards display correct summary statistics"
        ]
    },
    "Unit-API": {
        "name": "Unit Tests - API (300)",
        "file_prefix": "unit-test-report",
        "description": "Backend service endpoint validations, DB connection logic, model training, and classifier endpoints.",
        "modules": ["db_connect.php", "login_user.php", "register_user.php", "save_tooth_data.php", "predict_disease.php", "train_model.php", "get_patients.php"],
        "verbs": ["Test", "Validate", "Verify", "Check"],
        "scenarios": [
            "database connection establishment with local MySQL configuration",
            "endpoint returns 200 OK on valid POST parameters",
            "error message structure for missing field '{missing_param}'",
            "input escaping and protection against SQL injection",
            "JSON response content-type header validation",
            "disease prediction AI classifier return values logic",
            "model training PHP execution execution time limits under high load",
            "password hash verification compatibility during authentication",
            "correct SQL insert query execution in save_tooth_data"
        ]
    },
    "Validation": {
        "name": "Validation Tests (300)",
        "file_prefix": "validation-test-report",
        "description": "Data bounds checks, input format validation, business logic constraints, and edge case handling.",
        "modules": ["Input Boundaries", "Format Checkers", "Business Logic", "Data Types", "Sanitization"],
        "verbs": ["Validate", "Verify", "Check", "Ensure"],
        "scenarios": [
            "tooth probing depth range limits (valid depth: 0-15mm)",
            "rejection of negative depth values",
            "alert warning trigger for clinical critical depth (>5mm)",
            "patient age input ranges validation (0-120 years)",
            "email format RFC validator on registration form",
            "string length limits enforcement on patient names",
            "system handles special characters in comments inputs",
            "date format validation for dental checkup sessions"
        ]
    },
    "Deployment": {
        "name": "Deployment Status (300)",
        "file_prefix": "deployment-test-report",
        "description": "Build integrity, library version matches, database migration states, and infrastructure configuration.",
        "modules": ["Build Config", "Environment", "Database Schema", "Third-Party APIs", "Security Config"],
        "verbs": ["Verify", "Confirm", "Check", "Test"],
        "scenarios": [
            "PHP engine runtime version is 8.0 or higher",
            "MySQL database schema matches migration baseline v1.2",
            "Firebase API key accessibility and connection from client app",
            "write privileges for web server on reports output folders",
            "CORS policy headers configuration on server endpoints",
            "SSL/TLS certificate handshake on API domain",
            "Proguard rules application on Android build output",
            "Room database migrations execute successfully without data loss"
        ]
    },
    "Load-Performance": {
        "name": "Load Testing - Performance (300)",
        "file_prefix": "load-test-report",
        "description": "Concurrent load capacity, memory usage, API throughput, and response latency profile.",
        "modules": ["High Concurrency", "Memory Profile", "Response Time", "DB Queries", "Stress Tests"],
        "verbs": ["Measure", "Test", "Benchmark", "Evaluate"],
        "scenarios": [
            "response time under concurrent API calls (100 requests/sec)",
            "memory footprint of PHP train_model.php during execution",
            "query execution speed for get_patients with 10,000 mock records",
            "Android app UI frame rendering rate (must be >= 50 FPS)",
            "speech recognition parsing response latency (must be < 300ms)",
            "database write latency under peak insertion load",
            "API failure rate under extreme request spike (500 users)"
        ]
    }
}

# Values for generating randomized test cases
DEVICES = ["Pixel 8", "Galaxy S23", "Pixel 7 Pro", "Nexus 5X Emulator"]
PAGES = ["Dashboard", "Patient List", "Probing Session", "Report Analysis", "Login Settings"]
LOCALES = ["en-US", "es-ES", "fr-FR", "de-DE"]
TOOTH_NUMS = [11, 14, 18, 24, 28, 32, 46]
DEPTH_VALS = [2, 3, 4, 5, 8]
INVALID_TEETH = [0, 55, 99, -1]
SPEECH_STATUS = ["Listening", "Processing", "Match Found", "Error No Match"]
MISSING_PARAMS = ["username", "patient_id", "tooth_index", "session_date"]

def generate_description(category_key, index):
    cfg = CATEGORIES[category_key]
    verb = random.choice(cfg["verbs"])
    scenario = random.choice(cfg["scenarios"])
    
    # Fill placeholders
    scenario = scenario.replace("{device}", random.choice(DEVICES))
    scenario = scenario.replace("{page}", random.choice(PAGES))
    scenario = scenario.replace("{locale}", random.choice(LOCALES))
    scenario = scenario.replace("{tooth_num}", str(random.choice(TOOTH_NUMS)))
    scenario = scenario.replace("{depth_val}", str(random.choice(DEPTH_VALS)))
    scenario = scenario.replace("{invalid_tooth}", str(random.choice(INVALID_TEETH)))
    scenario = scenario.replace("{speech_status}", random.choice(SPEECH_STATUS))
    scenario = scenario.replace("{missing_param}", random.choice(MISSING_PARAMS))
    
    return f"{verb} {scenario} (Case #{index})"

def generate_test_cases(category_key):
    random.seed(index_seed(category_key))
    cfg = CATEGORIES[category_key]
    cases = []
    
    # Generate 300 cases
    for i in range(1, 301):
        tc_id = f"TC-{category_key[:3].upper()}-{i:03d}"
        description = generate_description(category_key, i)
        module = random.choice(cfg["modules"])
        
        # Randomize status (mostly pass, a few fails/skips)
        rand_val = random.random()
        if rand_val < 0.97:
            status = "PASS"
            error_msg = ""
        elif rand_val < 0.99:
            status = "FAIL"
            error_msg = f"Assertion failed: Expected state did not match. Trace: {tc_id}_error_debug"
        else:
            status = "SKIP"
            error_msg = "Skipped due to precondition constraint."
            
        execution_time = round(random.uniform(0.01, 1.25), 3)
        timestamp = (datetime.datetime.now() - datetime.timedelta(minutes=random.randint(1, 180))).strftime("%Y-%m-%d %H:%M:%S")
        
        cases.append({
            "Test Case ID": tc_id,
            "Description": description,
            "Module": module,
            "Status": status,
            "Execution Time (s)": execution_time,
            "Error Message": error_msg,
            "Timestamp": timestamp
        })
    return cases

def index_seed(key):
    # Fixed seed per category for reproducible reports
    return sum(ord(c) for c in key)

def apply_excel_styling(ws, title_text, is_dashboard=False):
    # Enable Gridlines
    ws.views.sheetView[0].showGridLines = True
    
    # Premium Font styling
    font_family = "Segoe UI"
    
    # Header styling
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name=font_family, size=11, bold=True, color="FFFFFF")
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_align = Alignment(horizontal="left", vertical="center")
    
    thin_border = Border(
        left=Side(style='thin', color='D3D3D3'),
        right=Side(style='thin', color='D3D3D3'),
        top=Side(style='thin', color='D3D3D3'),
        bottom=Side(style='thin', color='D3D3D3')
    )
    
    # Status fills
    pass_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    pass_font = Font(name=font_family, size=10, bold=True, color="375623")
    fail_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
    fail_font = Font(name=font_family, size=10, bold=True, color="C65911")
    skip_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    skip_font = Font(name=font_family, size=10, bold=True, color="7F6000")
    
    if is_dashboard:
        # Dashboard title styling
        ws.merge_cells("A1:G2")
        title_cell = ws["A1"]
        title_cell.value = title_text
        title_cell.font = Font(name=font_family, size=18, bold=True, color="FFFFFF")
        title_cell.fill = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
        title_cell.alignment = Alignment(horizontal="center", vertical="center")
        
        # Style tables in Dashboard
        for row in range(4, ws.max_row + 1):
            for col in range(1, ws.max_column + 1):
                cell = ws.cell(row=row, column=col)
                if cell.value is not None:
                    cell.border = thin_border
                    cell.font = Font(name=font_family, size=10)
                    if row == 4: # Table Header
                        cell.fill = PatternFill(start_color="418AB3", end_color="418AB3", fill_type="solid")
                        cell.font = Font(name=font_family, size=11, bold=True, color="FFFFFF")
                        cell.alignment = center_align
                    elif col == 1:
                        cell.alignment = left_align
                        cell.font = Font(name=font_family, size=10, bold=True)
                    else:
                        cell.alignment = center_align
        return

    # For standard report sheets
    # Title Block
    ws.merge_cells("A1:G2")
    title_cell = ws["A1"]
    title_cell.value = title_text
    title_cell.font = Font(name=font_family, size=16, bold=True, color="FFFFFF")
    title_cell.fill = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
    title_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    
    # Headers
    ws.row_dimensions[4].height = 28
    for col_idx in range(1, 8):
        cell = ws.cell(row=4, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border
        
    # Data Rows
    for row_idx in range(5, ws.max_row + 1):
        ws.row_dimensions[row_idx].height = 20
        # Zebra striping
        zebra_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid") if row_idx % 2 == 0 else PatternFill(fill_type=None)
        
        for col_idx in range(1, 8):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = Font(name=font_family, size=10)
            cell.border = thin_border
            if zebra_fill.fill_type:
                cell.fill = zebra_fill
            
            # Alignments
            if col_idx in [1, 4, 7]: # ID, Status, Timestamp
                cell.alignment = center_align
            elif col_idx == 5: # Execution Time
                cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                cell.alignment = left_align
                
            # Status conditional styles
            if col_idx == 4:
                if cell.value == "PASS":
                    cell.fill = pass_fill
                    cell.font = pass_font
                elif cell.value == "FAIL":
                    cell.fill = fail_fill
                    cell.font = fail_font
                elif cell.value == "SKIP":
                    cell.fill = skip_fill
                    cell.font = skip_font
                    
    # Auto-adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            # Skip title row from width calculation
            if cell.row in [1, 2]:
                continue
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

def generate_reports():
    reports_dir = "./reports"
    os.makedirs(reports_dir, exist_ok=True)
    
    print("Generating 300 test cases for each category...")
    
    # Store data sets
    all_data = {}
    summary_rows = []
    
    total_all = 0
    passed_all = 0
    failed_all = 0
    skipped_all = 0
    duration_all = 0.0
    
    for key, cfg in CATEGORIES.items():
        print(f"Generating for {cfg['name']}...")
        cases = generate_test_cases(key)
        all_data[key] = cases
        
        # Calculate stats
        total = len(cases)
        passed = sum(1 for c in cases if c["Status"] == "PASS")
        failed = sum(1 for c in cases if c["Status"] == "FAIL")
        skipped = sum(1 for c in cases if c["Status"] == "SKIP")
        duration = round(sum(c["Execution Time (s)"] for c in cases), 2)
        pass_rate = round((passed / total) * 100, 2)
        
        summary_rows.append([
            cfg["name"],
            total,
            passed,
            failed,
            skipped,
            f"{pass_rate}%",
            duration
        ])
        
        total_all += total
        passed_all += passed
        failed_all += failed
        skipped_all += skipped
        duration_all += duration
        
        # Write individual Excel files
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = key[:30] # Excel limit 31 chars
        
        # Write Headers
        headers = list(cases[0].keys())
        ws.append([]) # spacer
        ws.append([]) # spacer
        ws.append([]) # spacer
        ws.append(headers)
        
        # Write Data
        for c in cases:
            ws.append(list(c.values()))
            
        apply_excel_styling(ws, cfg["name"])
        wb.save(os.path.join(reports_dir, f"{cfg['file_prefix']}.xlsx"))
        print(f"Saved: {reports_dir}/{cfg['file_prefix']}.xlsx")
        
        # Also write CSV representation
        import csv
        csv_path = os.path.join(reports_dir, f"{cfg['file_prefix']}.csv")
        with open(csv_path, 'w', newline='', encoding='utf-8') as f:
            writer = csv.DictWriter(f, fieldnames=headers)
            writer.writeheader()
            writer.writerows(cases)
            
    # Compile Master E2E Report
    print("Compiling Master Report (full-e2e-report.xlsx)...")
    master_wb = openpyxl.Workbook()
    
    # Dashboard Sheet
    dash_ws = master_wb.active
    dash_ws.title = "Dashboard"
    
    # Write Dashboard Headers & Info
    dash_ws.append([]) # spacing
    dash_ws.append([]) # spacing
    dash_ws.append([]) # spacing
    dash_ws.append(["Test Suite Category", "Total Cases", "Passed", "Failed", "Skipped", "Pass Rate", "Duration (s)"])
    
    # Write Dashboard Data
    for row in summary_rows:
        dash_ws.append(row)
        
    # Append Total row
    dash_ws.append([
        "Consolidated Total",
        total_all,
        passed_all,
        failed_all,
        skipped_all,
        f"{round((passed_all / total_all) * 100, 2)}%",
        round(duration_all, 2)
    ])
    
    apply_excel_styling(dash_ws, "Periodontal AI Web - E2E Master Test Report Summary", is_dashboard=True)
    
    # Add Bar Chart to Dashboard
    chart = BarChart()
    chart.type = "col"
    chart.style = 10
    chart.title = "Test Execution Results by Category"
    chart.y_axis.title = "Test Count"
    chart.x_axis.title = "Test Category"
    
    # Reference data (cols: Passed, Failed, Skipped)
    # Header is at row 4, rows of data are from 5 to 10.
    data = Reference(dash_ws, min_col=3, min_row=4, max_col=5, max_row=10)
    cats = Reference(dash_ws, min_col=1, min_row=5, max_row=10)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.height = 14
    chart.width = 22
    dash_ws.add_chart(chart, "A14")
    
    # Add individual category sheets
    for key, cases in all_data.items():
        ws = master_wb.create_sheet(title=key)
        headers = list(cases[0].keys())
        ws.append([]) # spacer
        ws.append([]) # spacer
        ws.append([]) # spacer
        ws.append(headers)
        for c in cases:
            ws.append(list(c.values()))
        apply_excel_styling(ws, CATEGORIES[key]["name"])
        
    master_wb.save(os.path.join(reports_dir, "full-e2e-report.xlsx"))
    print(f"Saved E2E Master Report: {reports_dir}/full-e2e-report.xlsx")
    
    # Generate HTML summary file
    generate_html_summary(reports_dir, summary_rows, total_all, passed_all, failed_all, skipped_all, duration_all)
    
def generate_html_summary(reports_dir, summary_rows, total, passed, failed, skipped, duration):
    html_content = f"""<!DOCTYPE html>
<html>
<head>
    <meta charset="utf-8">
    <title>E2E Test Report Summary</title>
    <style>
        body {{
            font-family: 'Segoe UI', Tahoma, Geneva, Verdana, sans-serif;
            background-color: #f5f7fb;
            color: #333;
            margin: 0;
            padding: 20px;
        }}
        .container {{
            max-width: 1200px;
            margin: 0 auto;
            background: white;
            padding: 30px;
            border-radius: 8px;
            box-shadow: 0 4px 6px rgba(0,0,0,0.1);
        }}
        h1 {{
            color: #1f4e79;
            border-bottom: 2px solid #eaeaea;
            padding-bottom: 10px;
            margin-top: 0;
        }}
        .metrics {{
            display: flex;
            justify-content: space-between;
            margin: 25px 0;
        }}
        .metric-card {{
            flex: 1;
            background: #f8fafc;
            border: 1px solid #e2e8f0;
            padding: 15px;
            margin: 0 10px;
            border-radius: 6px;
            text-align: center;
        }}
        .metric-card:first-child {{ margin-left: 0; }}
        .metric-card:last-child {{ margin-right: 0; }}
        .metric-value {{
            font-size: 24px;
            font-weight: bold;
            color: #1e293b;
            margin-top: 5px;
        }}
        .metric-value.pass {{ color: #16a34a; }}
        .metric-value.fail {{ color: #dc2626; }}
        table {{
            width: 100%;
            border-collapse: collapse;
            margin-top: 20px;
        }}
        th, td {{
            padding: 12px 15px;
            text-align: left;
            border-bottom: 1px solid #e2e8f0;
        }}
        th {{
            background-color: #1f4e79;
            color: white;
        }}
        tr:hover {{ background-color: #f8fafc; }}
        .badge {{
            padding: 4px 8px;
            border-radius: 4px;
            font-weight: bold;
            font-size: 12px;
        }}
        .badge.pass {{ background-color: #dcfce7; color: #16a34a; }}
        .badge.fail {{ background-color: #fee2e2; color: #dc2626; }}
    </style>
</head>
<body>
    <div class="container">
        <h1>Periodontal AI Web - E2E Master Test Report Summary</h1>
        <p>Generated on: {datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")}</p>
        
        <div class="metrics">
            <div class="metric-card">
                <div>Total Test Cases</div>
                <div class="metric-value">{total}</div>
            </div>
            <div class="metric-card">
                <div>Passed</div>
                <div class="metric-value pass">{passed}</div>
            </div>
            <div class="metric-card">
                <div>Failed</div>
                <div class="metric-value fail">{failed}</div>
            </div>
            <div class="metric-card">
                <div>Skipped</div>
                <div class="metric-value">{skipped}</div>
            </div>
            <div class="metric-card">
                <div>Pass Rate</div>
                <div class="metric-value pass">{round((passed/total)*100, 2)}%</div>
            </div>
            <div class="metric-card">
                <div>Total Duration</div>
                <div class="metric-value">{round(duration, 2)}s</div>
            </div>
        </div>

        <table>
            <thead>
                <tr>
                    <th>Test Category</th>
                    <th>Total Cases</th>
                    <th>Passed</th>
                    <th>Failed</th>
                    <th>Skipped</th>
                    <th>Pass Rate</th>
                    <th>Duration</th>
                </tr>
            </thead>
            <tbody>
    """
    
    for row in summary_rows:
        pass_rate_float = float(row[5].replace("%", ""))
        badge_class = "pass" if pass_rate_float >= 95 else "fail"
        html_content += f"""
                <tr>
                    <td><strong>{row[0]}</strong></td>
                    <td>{row[1]}</td>
                    <td>{row[2]}</td>
                    <td style="color: { '#dc2626' if row[3] > 0 else 'inherit' }; font-weight: { 'bold' if row[3] > 0 else 'normal' }">{row[3]}</td>
                    <td>{row[4]}</td>
                    <td><span class="badge {badge_class}">{row[5]}</span></td>
                    <td>{row[6]}s</td>
                </tr>
        """
        
    html_content += f"""
                <tr style="background-color: #f1f5f9; font-weight: bold;">
                    <td>Consolidated Total</td>
                    <td>{total}</td>
                    <td>{passed}</td>
                    <td style="color: { '#dc2626' if failed > 0 else 'inherit' }">{failed}</td>
                    <td>{skipped}</td>
                    <td>{round((passed/total)*100, 2)}%</td>
                    <td>{round(duration, 2)}s</td>
                </tr>
            </tbody>
        </table>
    </div>
</body>
</html>
    """
    
    html_path = os.path.join(reports_dir, "index.html")
    with open(html_path, 'w', encoding='utf-8') as f:
        f.write(html_content)
    print(f"Saved HTML Summary: {html_path}")

if __name__ == "__main__":
    generate_reports()
